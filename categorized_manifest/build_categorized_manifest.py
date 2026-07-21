#!/usr/bin/env python3
"""Build JSON manifests + a status CSV from TO_25-203_All_Datasets_Categorized.xlsx.

The spreadsheet only stores a bare "File Name" per row (the person who filled it
in did not save the full dataset-relative path). This script re-resolves each
file to its real image + ground-truth location under /shared/data, records the
categorization tags, and reports what could and could not be found.

Datasets handled: TuSimple, BDD100K, CurveLane, CuLane.

Per-dataset resolution rules (derived from the on-disk layout):
  BDD100K   File Name = "<id>.jpg". Image under images/{val,train,test}.
            Lane GT = ll_seg_annotations/{val,train}/<id>.png  (NO test masks).
  CuLane    File Name = "<seg>.mp4" -> a *segment directory* <seg>.MP4 holding
            ~360 frames, each with a .jpg + per-frame .lines.txt GT. Segment
            lives under one of the driver_* dirs. laneseg masks live under
            laneseg_label_w16{,_test}/<driver>/<seg>.MP4/. No single frame is
            named in the sheet, so we resolve the segment dir (a video clip).
  CurveLane File Name = "<hash>.jpg". Image under {train,valid}/images.
            Lane GT = same-split labels/<hash>.lines.json. test/ has no labels.
  TuSimple  File Name = a bare numeric id (e.g. "985", "64500"). TuSimple images
            live at clips/<date>/<timestamp>/<frame>.jpg -- there is no numeric
            id in that layout and no mapping was saved, so these are UNRESOLVED.
"""
import csv
import json
import os
from pathlib import Path

import openpyxl

XLSX = Path(__file__).resolve().parent.parent / "TO_25-203_All_Datasets_Categorized.xlsx"
OUT = Path(__file__).resolve().parent / "output"
DATA = Path("/shared/data")

SHEETS = ["TuSimple", "BDD100K", "CurveLane", "CuLane"]

# Column layout (1-indexed): 1=#, 2=File Name, 3..8=summary text per dimension,
# 9..end=one-hot tag columns. Row 1=group headers, row 2=sub-headers.
SUMMARY_COLS = list(range(3, 9))   # 6 summary dimensions
TAG_START_COL = 9


def parse_sheet(ws):
    """Return (summary_dim_names, tag_names, rows) for one sheet."""
    header2 = [c.value for c in ws[2]]
    summary_dims = [header2[c - 1] for c in SUMMARY_COLS]
    tag_names = [header2[c - 1] for c in range(TAG_START_COL, ws.max_column + 1)]
    rows = []
    for r in range(3, ws.max_row + 1):
        idx = ws.cell(row=r, column=1).value
        fname = ws.cell(row=r, column=2).value
        if fname is None:
            continue
        fname = str(fname).strip()
        # skip the trailing "TOTAL (count)" summary row
        if idx is None and fname.upper().startswith("TOTAL"):
            continue
        summary = {}
        for dim, col in zip(summary_dims, SUMMARY_COLS):
            v = ws.cell(row=r, column=col).value
            if v is not None and str(v).strip():
                summary[dim] = str(v).strip()
        labels = []
        for i, col in enumerate(range(TAG_START_COL, ws.max_column + 1)):
            v = ws.cell(row=r, column=col).value
            if v is not None and str(v).strip() not in ("", "0"):
                labels.append(tag_names[i])
        rows.append({"row_index": idx, "file_name": fname,
                     "summary": summary, "labels": labels})
    return summary_dims, tag_names, rows


# ---------------------------------------------------------------------------
# Resolvers: each returns a dict of resolution fields merged into the sample.
# ---------------------------------------------------------------------------

def resolve_bdd(fname):
    bdd = DATA / "bdd100k"
    name = fname if fname.lower().endswith(".jpg") else fname + ".jpg"
    base = name[:-4]
    img_path, split = None, None
    for sp in ("val", "train", "test"):
        p = bdd / "images" / sp / name
        if p.exists():
            img_path, split = str(p), sp
            break
    gt_path = None
    if split in ("val", "train"):
        m = bdd / "ll_seg_annotations" / split / (base + ".png")
        if m.exists():
            gt_path = str(m)
    gt_note = "" if gt_path else (
        "test split has no ll_seg lane masks" if split == "test"
        else ("image not found" if img_path is None else "lane mask missing"))
    return {"image_path": img_path, "split": split,
            "ground_truth_path": gt_path, "gt_type": "lane_mask_png",
            "image_found": img_path is not None, "gt_found": gt_path is not None,
            "note": gt_note}


def _build_culane_index():
    idx = {}
    for drv in sorted((DATA / "culane").glob("driver_*")):
        if not drv.is_dir():
            continue
        for seg in drv.iterdir():
            if seg.is_dir():
                idx[seg.name.lower()] = (drv.name, seg)
    return idx


def resolve_culane(fname, index):
    # sheet stores "<seg>.mp4"; on disk the segment dir is "<seg>.MP4"
    seg_name = fname if fname.lower().endswith(".mp4") else fname + ".mp4"
    hit = index.get(seg_name.lower())
    if hit is None:
        return {"image_path": None, "split": None, "ground_truth_path": None,
                "gt_type": "culane_segment", "image_found": False,
                "gt_found": False, "num_frames": 0,
                "note": "segment directory not found under any driver_*"}
    driver, seg_dir = hit
    frames = sorted(seg_dir.glob("*.jpg"))
    lines = sorted(seg_dir.glob("*.lines.txt"))
    # laneseg masks: test drivers live under laneseg_label_w16_test
    gt_dir = None
    for base in ("laneseg_label_w16", "laneseg_label_w16_test"):
        cand = DATA / "culane" / base / driver / seg_dir.name
        if cand.is_dir() and any(cand.glob("*.png")):
            gt_dir = str(cand)
            break
    return {
        "image_path": str(seg_dir),           # segment dir (whole video clip)
        "representative_frame": str(frames[0]) if frames else None,
        "driver": driver,
        "split": "test" if "test" in (gt_dir or "") else "train/val",
        "num_frames": len(frames),
        "ground_truth_path": gt_dir,          # laneseg mask dir
        "lines_txt_count": len(lines),        # per-frame polyline GT co-located
        "gt_type": "culane_segment (per-frame .lines.txt + laneseg masks)",
        "image_found": len(frames) > 0,
        "gt_found": len(lines) > 0 or gt_dir is not None,
        "note": "sheet names a video segment, not a single frame"
                if frames else "segment dir has no frames",
    }


def resolve_curvelane(fname):
    cl = DATA / "Curvelanes"
    name = fname if fname.lower().endswith(".jpg") else fname + ".jpg"
    base = name[:-4]
    img_path, split, gt_path = None, None, None
    for sp in ("valid", "train", "test"):
        p = cl / sp / "images" / name
        if p.exists():
            img_path, split = str(p), sp
            g = cl / sp / "labels" / (base + ".lines.json")
            if g.exists():
                gt_path = str(g)
            break
    note = ""
    if img_path is None:
        note = "image not found in train/valid/test"
    elif gt_path is None:
        note = "test split has no labels" if split == "test" else "label json missing"
    return {"image_path": img_path, "split": split,
            "ground_truth_path": gt_path, "gt_type": "lines_json",
            "image_found": img_path is not None, "gt_found": gt_path is not None,
            "note": note}


def resolve_tusimple(fname):
    # Bare numeric id with no saved mapping to clips/<date>/<ts>/<frame>.jpg.
    return {"image_path": None, "split": None, "ground_truth_path": None,
            "gt_type": "tusimple_lane_json", "image_found": False,
            "gt_found": False,
            "note": "bare numeric id; no path/mapping saved -> UNRESOLVED "
                    "(needs the id->clip mapping from the person who labeled it)"}


def main():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    culane_index = _build_culane_index()
    OUT.mkdir(parents=True, exist_ok=True)

    csv_rows = []
    for sheet in SHEETS:
        ws = wb[sheet]
        summary_dims, tag_names, rows = parse_sheet(ws)
        dataset = sheet.lower()
        samples = []
        for row in rows:
            fn = row["file_name"]
            if sheet == "BDD100K":
                res = resolve_bdd(fn)
            elif sheet == "CuLane":
                res = resolve_culane(fn, culane_index)
            elif sheet == "CurveLane":
                res = resolve_curvelane(fn)
            else:
                res = resolve_tusimple(fn)
            sample = {
                "sample_id": f"{dataset}_{row['row_index']}",
                "file_name": fn,
                "tags": {"summary": row["summary"], "labels": row["labels"]},
                **res,
            }
            samples.append(sample)
            csv_rows.append({
                "dataset": sheet,
                "row": row["row_index"],
                "file_name": fn,
                "split": res.get("split") or "",
                "image_found": res["image_found"],
                "gt_found": res["gt_found"],
                "image_path": res.get("image_path") or "",
                "ground_truth_path": res.get("ground_truth_path") or "",
                "num_labels": len(row["labels"]),
                "note": res.get("note", ""),
            })

        n_img = sum(s["image_found"] for s in samples)
        n_gt = sum(s["gt_found"] for s in samples)
        manifest = {
            "metadata": {
                "dataset": dataset,
                "source_xlsx": XLSX.name,
                "source_sheet": sheet,
                "num_samples": len(samples),
                "num_image_found": n_img,
                "num_gt_found": n_gt,
                "num_image_missing": len(samples) - n_img,
                "num_gt_missing": len(samples) - n_gt,
                "summary_dimensions": summary_dims,
                "tag_vocabulary": tag_names,
            },
            "samples": samples,
        }
        out_json = OUT / f"manifest_{dataset}.json"
        out_json.write_text(json.dumps(manifest, indent=2))
        print(f"{sheet:10s} samples={len(samples):3d}  "
              f"image_found={n_img:3d}  gt_found={n_gt:3d}  -> {out_json.name}")

    csv_path = OUT / "resolution_status.csv"
    with csv_path.open("w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(csv_rows[0].keys()))
        w.writeheader()
        w.writerows(csv_rows)
    print(f"\nStatus CSV -> {csv_path}  ({len(csv_rows)} rows)")


if __name__ == "__main__":
    main()
